from openpyxl import load_workbook
import db_config_1 as db

# Excel file path
excel_file = "D:\Julee\CODE\Metro_cash_and_carry\pincode_sheet_data.xlsx"
# Connect to MySQL
connection = db.get_db_connection(db.db_name)
try:
    # Load workbook and specific sheet
    wb = load_workbook(excel_file)
    # sheet = wb["pincode"]   # <-- sheet name
    sheet = wb["new_pincode"]   # <-- sheet name
    with connection.cursor() as cursor:
        # Insert query
        insert_query = f"""
        INSERT INTO {db.pincode_table} (city, pincode_area_name, pincode, state, state_code, status)
        VALUES (%s, %s, %s, %s, %s, %s)
        """
        data = []
        # Skip header row (start from row 2)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            city = row[0]
            area = row[1]
            pincode = str(row[2])
            state = row[3]
            state_code = row[4]
            data.append((city, area, pincode, state, state_code, "pending"))
        # Bulk insert
        cursor.executemany(insert_query, data)
    connection.commit()
finally:
    connection.close()