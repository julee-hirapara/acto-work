# # import json
# # from datetime import datetime
# # import db_config_1 as db
# # from openpyxl import Workbook
# # import os
# # from openpyxl.styles import Font
# #
# # # ----------------------------
# # # 📥 FETCH DATA FROM DB
# # # ----------------------------
# # def export_all():
# #     conn = db.get_db_connection(db.db_name)
# #     cursor = conn.cursor()
# #     cursor.execute(f"SELECT DISTINCT pincode, city FROM {db.product_link_table} where pincode='700001' limit 1")
# #     rows = cursor.fetchall()
# #     cursor.close()
# #     conn.close()
# #     for r in rows:
# #         print(f"\n🚀 Exporting: {r['city']} ({r['pincode']})")
# #         export_products(r["pincode"], r["city"])
# #
# #
# # def fetch_products(pincode, city):
# #     conn = db.get_db_connection(db.db_name)
# #     cursor = conn.cursor()
# #
# #     query = f"""
# #     SELECT * FROM {db.product_link_table}
# #     WHERE pincode = %s AND city = %s limit 1
# #     """
# #     cursor.execute(query, (pincode, city))
# #     rows = cursor.fetchall()
# #     cursor.close()
# #     conn.close()
# #     print(f"📦 Fetched {len(rows)} rows from DB")
# #     return rows
# #
# #
# # # ----------------------------
# # # 🔄 TRANSFORM DATA
# # # ----------------------------
# # def transform_data(rows):
# #     output = []
# #     for row in rows:
# #         try:
# #             # 🔥 Convert images JSON string → list
# #             images = json.loads(row["images"]) if row["images"] else []
# #             # ✅ Take first image only
# #             first_image = images[0] if images else ""
# #             full_path = row.get("pagesave")
# #             pdp_data = row.get("pdp_data")
# #             print(pdp_data)
# #             file_name = os.path.basename(full_path)
# #             created_at = row.get("created_at")
# #             if created_at:
# #                 try:
# #                     # convert string → datetime
# #                     dt = datetime.strptime(str(created_at), "%Y-%m-%d %H:%M:%S")
# #                     # format to required format with milliseconds
# #                     formatted_dt = dt.strftime("%Y-%m-%dT%H:%M:%S") + ":000"
# #                 except:
# #                     formatted_dt = created_at
# #             else:
# #                 formatted_dt = None
# #
# #             # "scrap_date_time": formatted_dt,
# #             output.append({
# #                 "item_code": row.get("item_code"),
# #                 "image": first_image,
# #                 "name": row.get("product_name"),
# #                 "brand": row.get("brand_name"),
# #                 "weight": row.get("weight"),
# #                 "net_quantity": row.get("net_quantity"),
# #                 "country_of_origin": row.get("country_of_origin"),
# #                 "store_id": row.get("store_id"),
# #                 "mrp": row.get("mrp"),
# #                 "rrp": row.get("rrp"),
# #                 "discount": row.get("discount"),
# #                 "in_stock": row.get("in_stock"),
# #                 "minimum_quantity": row.get("minimum_quantity"),
# #                 "maximum_quantity": row.get("maximum_quantity"),
# #                 "city": row.get("city"),
# #                 "pincode": row.get("pincode"),
# #                 "scrap_date_time": formatted_dt,
# #                 "pagesave": file_name
# #             })
# #         except Exception as e:
# #             print("⚠️ Error processing row:", e)
# #     return output
# #
# #
# # # ----------------------------
# # # 💾 SAVE EXCEL FILE
# # # ----------------------------
# # def save_to_excel(data, pincode, city):
# #     file_path = os.path.join(db.RESULT_PATH, f"{city}_{pincode}.xlsx")
# #     wb = Workbook()
# #     ws = wb.active
# #     ws.title = "Products"
# #     headers = [
# #         "item_code", "image", "name", "brand", "weight",
# #         "net_quantity", "country_of_origin", "store_id",
# #         "mrp", "retail_price", "discount", "in_stock", "minimum_quantity", "maximum_quantity", "city", "pincode", "scrap_date_time", "pagesave"
# #     ]
# #     bold_font = Font(bold=True)
# #     ws.append(headers)
# #     for col_num in range(1, len(headers) + 1):
# #         ws.cell(row=1, column=col_num).font = bold_font
# #
# #     # # ✅ Write header
# #     # ws.append(headers)
# #     # ✅ Write data
# #     for row in data:
# #         # ws.append([
# #         #     (
# #         #         " ".join(str(val[0]).split()) if isinstance(val := row.get(col), list) and val else
# #         #         (" ".join(str(val).split()) if val and str(val).strip() else "N/A")
# #         #     )
# #         #     for col in headers
# #         # ])
# #         ws.append([
# #             (
# #                 " ".join(str(val[0]).split()) if isinstance(val := row.get("rrp" if col == "retail_price" else col),
# #                                                             list) and val else
# #                 (" ".join(str(val).split()) if val and str(val).strip() else "N/A")
# #             )
# #             for col in headers
# #         ])
# #         # ws.append([
# #         #     row.get("item_code"),
# #         #     row.get("image"),
# #         #     row.get("name"),
# #         #     row.get("brand"),
# #         #     row.get("weight"),
# #         #     row.get("net_quantity"),
# #         #     row.get("country_of_origin"),
# #         #     row.get("store_id"),
# #         #     row.get("mrp"),
# #         #     row.get("rrp"),
# #         #     row.get("discount"),
# #         #     row.get("in_stock"),
# #         #     row.get("city"),
# #         #     row.get("pincode")
# #         # ])
# #     # ws.append([
# #     #     (val if (val := row.get(col)) and str(val).strip() else "N/A")
# #     #     for col in headers
# #     # ])
# #     wb.save(file_path)
# #     print(f"✅ Excel file saved: {file_path}")
# #
# #
# # # ----------------------------
# # # 🚀 MAIN FUNCTION
# # # ----------------------------
# # def export_products(pincode, city):
# #     rows = fetch_products(pincode, city)
# #     if not rows:
# #         print("❌ No data found")
# #         return
# #     data = transform_data(rows)
# #     print(data)
# #     # save_to_excel(data, pincode, city)
# #
# #
# # # ----------------------------
# # # ▶ RUN
# # # ----------------------------
# # if __name__ == "__main__":
# #     export_all()
#
#

import json
import re
from datetime import datetime
import db_config_1 as db
from openpyxl import Workbook
from openpyxl.styles import Font
import os
import html


def clean_text(val):
    if not val:
        return None
    # 🔥 Decode HTML entities (&amp; → &)
    val = html.unescape(str(val))
    # 🔥 Remove extra spaces
    return " ".join(val.split())


# =========================
# CLEAN HELPERS
# =========================

# def clean_text(val):
#     if not val:
#         return None
#     return " ".join(str(val).split())


def clean_price(val):
    if not val:
        return None
    return re.sub(r"[^\d.]", "", str(val))


def clean_dict(d):
    if not isinstance(d, dict):
        return {}
    return {
        k: clean_text(v) if isinstance(v, str) else v
        for k, v in d.items()
    }


def clean_list_of_dict(lst):
    try:
        return [clean_dict(i) for i in lst if isinstance(i, dict)]
    except:
        return []


# =========================
# FETCH DATA FROM DB
# =========================

def export_all():
    conn = db.get_db_connection(db.db_name)
    cursor = conn.cursor()

    cursor.execute(f"""
        SELECT DISTINCT pincode, city
        FROM {db.product_link_table}
        WHERE pincode='700001'
        LIMIT 1
    """)

    rows = cursor.fetchall()
    cursor.close()
    conn.close()

    for r in rows:
        print(f"\n🚀 Exporting: {r['city']} ({r['pincode']})")
        export_products(r["pincode"], r["city"])


def fetch_products(pincode, city):
    conn = db.get_db_connection(db.db_name)
    cursor = conn.cursor()

    query = f"""
    SELECT * FROM {db.product_link_table}
    WHERE pincode = %s AND city = %s
    """

    cursor.execute(query, (pincode, city))
    rows = cursor.fetchall()

    cursor.close()
    conn.close()

    print(f"📦 Fetched {len(rows)} rows from DB")
    return rows


from datetime import datetime
import re


def parse_delivery_date(delivery_text):
    try:
        if not delivery_text:
            return None
        # extract date part like "Sat, 18 Apr"
        match = re.search(r"\w{3},\s*(\d{1,2}\s+\w{3})", delivery_text)
        if not match:
            return None
        date_str = match.group(1)  # "18 Apr"
        # add current year
        current_year = datetime.now().year
        full_date = f"{date_str} {current_year}"  # "18 Apr 2026"
        # convert to datetime
        dt = datetime.strptime(full_date, "%d %b %Y")
        # return in DD-MM-YYYY
        return dt.strftime("%d-%m-%Y")
    except Exception:
        return None


# =========================
# TRANSFORM DATA
# =========================

def transform_data(rows):
    output = []

    for row in rows:
        try:
            # ---------------------
            # PDP JSON LOAD
            # ---------------------
            pdp_json = json.loads(row["pdp_data"]) if row.get("pdp_data") else {}

            # ---------------------
            # BASIC INFO (PDP priority)
            # ---------------------
            # name = row.get("product_name")
            name = clean_text(pdp_json.get("product_name") or row.get("product_name"))
            brand = clean_text(pdp_json.get("brand_name") or row.get("brand_name"))

            # ---------------------
            # STOCK + SELLER
            # ---------------------
            stock = clean_text(pdp_json.get("stock"))
            if stock == 'in_stock':
                stock = "In Stock"
            else:
                stock = "Sold Out"
            seller = clean_text(pdp_json.get("seller"))
            # ---------------------
            # DESCRIPTION
            # ---------------------
            description = clean_text(pdp_json.get("description"))
            # ---------------------
            # IMAGE
            # ---------------------
            # ---------------------
            # IMAGE (FINAL FIX)
            # ---------------------
            images_list = pdp_json.get("images")
            if not images_list:
                raw_db_images = row.get("images")
                if raw_db_images:
                    try:
                        # try JSON load first
                        images_list = json.loads(raw_db_images)
                    except:
                        try:
                            # fallback: fix single quotes → double quotes
                            images_list = json.loads(raw_db_images.replace("'", '"'))
                        except:
                            # last fallback: treat as single string
                            images_list = [raw_db_images]
                else:
                    images_list = "N/A"
            # final safety
            if not images_list:
                images_list = "N/A"

            # if images_list == []:
            #     images_list = "N/A"
            # image = images_list[0] if images_list else None

            # ---------------------
            # PRICE
            # ---------------------
            mrp = clean_price(row.get("mrp"))
            retail_price = row.get("rrp")
            discount = row.get("discount")
            row_url = row.get("product_slug")
            pdp_url = "https://www.metrorrl.com/product/" + row_url
            # ---------------------
            # PROMOTIONS + TIERS (STRUCTURED)
            # ---------------------
            promotions = clean_list_of_dict(pdp_json.get("promotions", []))
            if promotions == []:
                promotions = "N/A"
            # tiers = clean_list_of_dict(pdp_json.get("quantity_price_tiers", []))
            # if tiers == []:
            #     tiers = "N/A"
            raw_tiers = pdp_json.get("quantity_price_tiers", [])
            tiers = clean_list_of_dict(raw_tiers) if raw_tiers else []
            # 🔥 ADD THIS
            min_qty = None
            max_qty = None
            if tiers:
                try:
                    min_qty = min(int(t.get("min_qty", 0)) for t in tiers if t.get("min_qty"))
                    max_qty = max(int(t.get("max_qty", 0)) for t in tiers if t.get("max_qty"))
                except:
                    min_qty = None
                    max_qty = None
            # keep your old behavior
            if tiers == []:
                tiers = "N/A"
            # ---------------------
            # DATE FORMAT
            # ---------------------
            delivery_text = pdp_json.get("delivery", None)
            delivery = parse_delivery_date(delivery_text)
            print("delivery ::::::::::: ", delivery)
            created_at = row.get("created_at")
            if created_at:
                try:
                    dt = datetime.strptime(str(created_at), "%Y-%m-%d %H:%M:%S")
                    formatted_dt = dt.strftime("%Y-%m-%dT%H:%M:%S") + ":000"
                except:
                    formatted_dt = created_at
            else:
                formatted_dt = None
            # 🔥 ADD THIS
            if stock == "Sold Out":
                mrp = "N/A"
                retail_price = "N/A"
                discount = "N/A"
                delivery = "N/A"
                promotions = "N/A"
            # ---------------------
            # FILE NAME
            # ---------------------
            full_path = row.get("pagesave")
            file_name = os.path.basename(full_path) if full_path else None
            # ---------------------
            # FINAL OUTPUT
            # ---------------------
            net_quantity = row.get("net_quantity").replace(" ", "")

            output.append({
                "item_code": row.get("item_code"),
                "Images": images_list,
                "name": name,
                "brand": brand,
                "weight": row.get("weight"),
                "net_quantity": net_quantity,
                "country_of_origin": row.get("country_of_origin"),
                "store_id": row.get("store_id"),
                # # "mrp": mrp,
                # "retail_price": row.get("rrp"),
                # "discount": row.get("discount"),
                "mrp": mrp,
                "retail_price": retail_price,
                "discount": discount,
                "stock": stock,
                "seller": seller,
                "description": description,
                "arrival_date": delivery,
                "promotions": promotions,
                "price_quantity_tiers": tiers,
                # "minimum_quantity": row.get("minimum_quantity"),
                # "maximum_quantity": row.get("maximum_quantity"),
                "minimum_quantity": min_qty,
                "maximum_quantity": max_qty,
                "city": row.get("city"),
                "pincode": row.get("pincode"),
                "url": pdp_url,
                "scrap_date_time": formatted_dt,
                "pagesave": file_name
            })

        except Exception as e:
            print("⚠️ Error processing row:", e)

    return output


# =========================
# SAVE TO EXCEL
# =========================

def save_to_excel(data, pincode, city):
    file_path = os.path.join(db.RESULT_PATH, f"{city}_{pincode}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"

    headers = [
        "Sr. No",
        "item_code", "Images", "name", "brand", "weight",
        "net_quantity", "country_of_origin", "store_id",
        "mrp", "retail_price", "discount", "stock",
        "seller", "description", "arrival_date", "promotions", "price_quantity_tiers",
        "minimum_quantity", "maximum_quantity",
        "city", "pincode", "url", "scrap_date_time", "pagesave"
    ]

    # Header bold
    bold_font = Font(bold=True)
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        ws.cell(row=1, column=col_num).font = bold_font

    for idx, row in enumerate(data, start=1):  # 👈 auto counter
        ws.append([
            idx,  # 👈 Sr.no value
            *[
                json.dumps(row.get(col), ensure_ascii=False)
                if isinstance(row.get(col), (list, dict))
                else clean_text(row.get(col)) if row.get(col)
                else "N/A"
                for col in headers[1:]  # 👈 skip Sr.no
            ]
        ])
    # # Data rows
    # for row in data:
    #     ws.append([
    #         json.dumps(row.get(col), ensure_ascii=False)
    #         if isinstance(row.get(col), (list, dict))
    #         else clean_text(row.get(col)) if row.get(col)
    #         else "N/A"
    #         for col in headers
    #     ])

    wb.save(file_path)
    print(f"✅ Excel saved: {file_path}")


# =========================
# MAIN FLOW
# =========================

def export_products(pincode, city):
    rows = fetch_products(pincode, city)

    if not rows:
        print("❌ No data found")
        return

    data = transform_data(rows)
    print(json.dumps(data, indent=2, ensure_ascii=False))

    save_to_excel(data, pincode, city)


# =========================
# RUN
# =========================

if __name__ == "__main__":
    export_all()

#
# import json
# import re
# from datetime import datetime
# import db_config_1 as db
# from openpyxl import Workbook
# from openpyxl.styles import Font
# import os
# import html
#
#
# # =========================
# # CLEAN HELPERS
# # =========================
#
# def clean_text(val):
#     if not val:
#         return None
#     val = html.unescape(str(val))
#     return " ".join(val.split())
#
#
# def clean_price(val):
#     if not val:
#         return None
#     return re.sub(r"[^\d.]", "", str(val))
#
#
# def clean_dict(d):
#     if not isinstance(d, dict):
#         return {}
#     return {
#         k: clean_text(v) if isinstance(v, str) else v
#         for k, v in d.items()
#     }
#
#
# def clean_list_of_dict(lst):
#     try:
#         return [clean_dict(i) for i in lst if isinstance(i, dict)]
#     except:
#         return []
#
#
# # =========================
# # FETCH DATA FROM DB
# # =========================
#
# def export_all():
#     conn = db.get_db_connection(db.db_name)
#     cursor = conn.cursor()
#
#     cursor.execute(f"""
#         SELECT DISTINCT pincode, city
#         FROM {db.product_link_table}
#         WHERE pincode='380001'
#         LIMIT 1
#     """)
#
#     rows = cursor.fetchall()
#     cursor.close()
#     conn.close()
#
#     for r in rows:
#         print(f"\n🚀 Exporting: {r['city']} ({r['pincode']})")
#         export_products(r["pincode"], r["city"])
#
#
# def fetch_products(pincode, city):
#     conn = db.get_db_connection(db.db_name)
#     cursor = conn.cursor()
#
#     query = f"""
#     SELECT * FROM {db.product_link_table}
#     WHERE pincode = %s AND city = %s
#     """
#
#     cursor.execute(query, (pincode, city))
#     rows = cursor.fetchall()
#
#     cursor.close()
#     conn.close()
#
#     print(f"📦 Fetched {len(rows)} rows from DB")
#     return rows
#
#
# # =========================
# # DATE PARSER
# # =========================
#
# def parse_delivery_date(delivery_text):
#     try:
#         if not delivery_text:
#             return None
#
#         match = re.search(r"\w{3},\s*(\d{1,2}\s+\w{3})", delivery_text)
#         if not match:
#             return None
#
#         date_str = match.group(1)
#         current_year = datetime.now().year
#         full_date = f"{date_str} {current_year}"
#
#         dt = datetime.strptime(full_date, "%d %b %Y")
#         return dt.strftime("%d-%m-%Y")
#
#     except Exception:
#         return None
#
#
# # =========================
# # TRANSFORM DATA
# # =========================
#
# def transform_data(rows):
#     output = []
#
#     for row in rows:
#         try:
#             pdp_json = json.loads(row["pdp_data"]) if row.get("pdp_data") else {}
#
#             # ---------------------
#             # BASIC INFO
#             # ---------------------
#             name = clean_text(pdp_json.get("product_name") or row.get("product_name"))
#             brand = clean_text(pdp_json.get("brand_name") or row.get("brand_name"))
#
#             # ---------------------
#             # STOCK
#             # ---------------------
#             stock_raw = clean_text(pdp_json.get("stock"))
#
#             if stock_raw == "in_stock":
#                 stock = "In Stock"
#             else:
#                 stock = "Sold Out"
#
#             # ---------------------
#             # SELLER + DESC
#             # ---------------------
#             seller = clean_text(pdp_json.get("seller"))
#             description = clean_text(pdp_json.get("description"))
#
#             # ---------------------
#             # IMAGES
#             # ---------------------
#             images_list = pdp_json.get("images", [])
#             if not images_list:
#                 images_list = []
#
#             # ---------------------
#             # PRICE
#             # ---------------------
#             mrp = clean_price(pdp_json.get("mrp") or row.get("mrp"))
#             retail_price = row.get("rrp")
#             discount = row.get("discount")
#
#             # ---------------------
#             # TIERS (IMPORTANT)
#             # ---------------------
#             raw_tiers = pdp_json.get("quantity_price_tiers", [])
#             tiers = clean_list_of_dict(raw_tiers) if raw_tiers else []
#
#             min_qty = None
#             max_qty = None
#
#             if tiers:
#                 try:
#                     min_qty = min(int(t.get("min_qty", 0)) for t in tiers if t.get("min_qty"))
#                     max_qty = max(int(t.get("max_qty", 0)) for t in tiers if t.get("max_qty"))
#                 except:
#                     min_qty = None
#                     max_qty = None
#
#             # ---------------------
#             # SOLD OUT RULE
#             # ---------------------
#             if stock == "Sold Out":
#                 mrp = "N/A"
#                 retail_price = "N/A"
#                 discount = "N/A"
#
#             # ---------------------
#             # PROMOTIONS
#             # ---------------------
#             promotions = clean_list_of_dict(pdp_json.get("promotions", []))
#
#             # ---------------------
#             # DELIVERY DATE
#             # ---------------------
#             delivery = parse_delivery_date(pdp_json.get("delivery"))
#
#             # ---------------------
#             # CREATED DATE
#             # ---------------------
#             created_at = row.get("created_at")
#             if created_at:
#                 try:
#                     dt = datetime.strptime(str(created_at), "%Y-%m-%d %H:%M:%S")
#                     formatted_dt = dt.strftime("%Y-%m-%dT%H:%M:%S") + ":000"
#                 except:
#                     formatted_dt = created_at
#             else:
#                 formatted_dt = None
#
#             # ---------------------
#             # FILE NAME
#             # ---------------------
#             file_name = os.path.basename(row.get("pagesave")) if row.get("pagesave") else None
#
#             # ---------------------
#             # FINAL OUTPUT
#             # ---------------------
#             output.append({
#                 "item_code": row.get("item_code"),
#                 "images": images_list,
#                 "name": name,
#                 "brand": brand,
#                 "weight": row.get("weight"),
#                 "net_quantity": row.get("net_quantity"),
#                 "country_of_origin": row.get("country_of_origin"),
#                 "store_id": row.get("store_id"),
#                 "mrp": mrp,
#                 "retail_price": retail_price,
#                 "discount": discount,
#                 "stock": stock,
#                 "seller": seller,
#                 "description": description,
#                 "arrival_date": delivery,
#                 "promotions": promotions,
#                 "price_quantity_tiers": tiers,
#                 "minimum_quantity": min_qty,
#                 "maximum_quantity": max_qty,
#                 "city": row.get("city"),
#                 "pincode": row.get("pincode"),
#                 "scrap_date_time": formatted_dt,
#                 "pagesave": file_name
#             })
#
#         except Exception as e:
#             print("⚠️ Error processing row:", e)
#
#     return output
#
#
# # =========================
# # SAVE TO EXCEL
# # =========================
#
# def save_to_excel(data, pincode, city):
#     file_path = os.path.join(db.RESULT_PATH, f"{city}_{pincode}.xlsx")
#
#     wb = Workbook()
#     ws = wb.active
#     ws.title = "Products"
#
#     headers = [
#         "Sr. No",
#         "item_code", "images", "name", "brand", "weight",
#         "net_quantity", "country_of_origin", "store_id",
#         "mrp", "retail_price", "discount", "stock",
#         "seller", "description", "arrival_date",
#         "promotions", "price_quantity_tiers",
#         "minimum_quantity", "maximum_quantity",
#         "city", "pincode", "scrap_date_time", "pagesave"
#     ]
#
#     bold_font = Font(bold=True)
#     ws.append(headers)
#
#     for col_num in range(1, len(headers) + 1):
#         ws.cell(row=1, column=col_num).font = bold_font
#
#     for idx, row in enumerate(data, start=1):
#         ws.append([
#             idx,
#             *[
#                 json.dumps(row.get(col), ensure_ascii=False)
#                 if isinstance(row.get(col), (list, dict))
#                 else clean_text(row.get(col)) if row.get(col)
#                 else "N/A"
#                 for col in headers[1:]
#             ]
#         ])
#
#     wb.save(file_path)
#     print(f"✅ Excel saved: {file_path}")
#
#
# # =========================
# # MAIN FLOW
# # =========================
#
# def export_products(pincode, city):
#     rows = fetch_products(pincode, city)
#
#     if not rows:
#         print("❌ No data found")
#         return
#
#     data = transform_data(rows)
#     print(json.dumps(data, indent=2, ensure_ascii=False))
#
#     save_to_excel(data, pincode, city)
#
#
# # =========================
# # RUN
# # =========================
#
# if __name__ == "__main__":
#     export_all()